#!/usr/bin/env python3
"""Export project table workbooks to JSON.

This is a small Godot-oriented exporter inspired by the team's TABLE SCRIPT
layout. It intentionally supports only the subset this project needs first:
#data sheets, primitive types, arrays, JSON cells, key validation, and skips.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install it with `python -m pip install openpyxl`."
    ) from exc


HEADER_KIND_ROW = 2
FIELD_NAME_ROW = 2
FIELD_TYPE_ROW = 3
EXPORT_FLAG_ROW = 4
DATA_START_ROW = 5
FIRST_FIELD_COLUMN = 2

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
SUPPORTED_TYPES = {"string", "int", "float", "bool", "json"}


@dataclass(frozen=True)
class ColumnSchema:
    index: int
    name: str
    data_type: str
    export: bool
    is_key: bool
    nullable: bool
    is_array: bool


@dataclass
class ExportError:
    workbook: Path
    sheet: str
    row: int
    column: int
    message: str

    def format(self) -> str:
        location = f"{self.workbook.name}:{self.sheet}!R{self.row}C{self.column}"
        return f"{location} - {self.message}"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Excel #data sheets to JSON.")
    parser.add_argument(
        "--input-dir",
        default="tables/excel",
        type=Path,
        help="Directory containing .xlsx/.xlsm table files.",
    )
    parser.add_argument(
        "--output-dir",
        default="data/generated",
        type=Path,
        help="Directory where JSON files are written.",
    )
    parser.add_argument(
        "--clean",
        action="store_true",
        help="Delete existing .json files in the output directory before exporting.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_dir = args.input_dir
    output_dir = args.output_dir

    if not input_dir.exists():
        print(f"Input directory does not exist: {input_dir}", file=sys.stderr)
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)
    if args.clean:
        for path in output_dir.glob("*.json"):
            path.unlink()

    errors: list[ExportError] = []
    exported_count = 0

    for workbook_path in sorted(iter_workbooks(input_dir)):
        exported_count += export_workbook(workbook_path, output_dir, errors)

    if errors:
        print("Export failed:")
        for error in errors:
            print(f"  {error.format()}")
        return 1

    print(f"Exported {exported_count} table(s) to {output_dir}")
    return 0


def iter_workbooks(input_dir: Path) -> list[Path]:
    workbooks: list[Path] = []
    for path in input_dir.rglob("*"):
        if path.name.startswith("~$"):
            continue
        if path.suffix.lower() in SUPPORTED_EXTENSIONS:
            workbooks.append(path)
    return workbooks


def export_workbook(
    workbook_path: Path, output_dir: Path, errors: list[ExportError]
) -> int:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    exported_count = 0

    for sheet in workbook.worksheets:
        if sheet.max_row < EXPORT_FLAG_ROW:
            continue
        if normalize_text(sheet.cell(HEADER_KIND_ROW, 1).value).lower() != "#data":
            continue

        schema = parse_schema(workbook_path, sheet, errors)
        if not schema:
            continue

        rows = parse_rows(workbook_path, sheet, schema, errors)
        if errors:
            continue

        output_path = output_dir / f"{sheet.title}.json"
        output_path.write_text(
            json.dumps(rows, ensure_ascii=False, indent="\t") + "\n",
            encoding="utf-8",
        )
        exported_count += 1

    return exported_count


def parse_schema(
    workbook_path: Path, sheet: Any, errors: list[ExportError]
) -> list[ColumnSchema]:
    schema: list[ColumnSchema] = []
    seen_names: set[str] = set()

    for column in range(FIRST_FIELD_COLUMN, sheet.max_column + 1):
        raw_name = sheet.cell(FIELD_NAME_ROW, column).value
        raw_type = sheet.cell(FIELD_TYPE_ROW, column).value
        raw_export = sheet.cell(EXPORT_FLAG_ROW, column).value

        if raw_name is None and raw_type is None and raw_export is None:
            continue

        name = normalize_text(raw_name)
        type_options = split_options(raw_type)
        export_flag = normalize_text(raw_export).lower()

        if not name:
            errors.append(schema_error(workbook_path, sheet, column, "Field name is empty."))
            continue
        if name in seen_names:
            errors.append(schema_error(workbook_path, sheet, column, f"Duplicate field: {name}"))
            continue
        seen_names.add(name)

        if not type_options:
            errors.append(schema_error(workbook_path, sheet, column, "Field type is empty."))
            continue

        is_array = name.endswith("[]")
        if is_array:
            name = name[:-2]

        data_type = ""
        is_key = False
        nullable = False
        for option in type_options:
            option_lower = option.lower()
            if option_lower == "key":
                is_key = True
            elif option_lower == "null":
                nullable = True
            elif option_lower.endswith("[]"):
                data_type = option_lower[:-2]
                is_array = True
            elif option_lower in SUPPORTED_TYPES:
                data_type = option_lower

        if not data_type:
            errors.append(
                schema_error(
                    workbook_path,
                    sheet,
                    column,
                    f"Unsupported field type/options: {normalize_text(raw_type)}",
                )
            )
            continue

        schema.append(
            ColumnSchema(
                index=column,
                name=name,
                data_type=data_type,
                export=export_flag == "data",
                is_key=is_key,
                nullable=nullable,
                is_array=is_array,
            )
        )

    if not any(column.export for column in schema):
        errors.append(schema_error(workbook_path, sheet, 1, "No exported fields marked as data."))
    if not any(column.is_key for column in schema):
        errors.append(schema_error(workbook_path, sheet, 1, "No key field is defined."))

    return schema


def parse_rows(
    workbook_path: Path,
    sheet: Any,
    schema: list[ColumnSchema],
    errors: list[ExportError],
) -> list[dict[str, Any]]:
    key_column = next(column for column in schema if column.is_key)
    exported_columns = [column for column in schema if column.export]
    rows: list[dict[str, Any]] = []
    seen_keys: set[str] = set()

    for row in range(DATA_START_ROW, sheet.max_row + 1):
        marker = normalize_text(sheet.cell(row, 1).value).lower()
        if marker == "skip":
            continue

        key_text = normalize_text(sheet.cell(row, key_column.index).value)
        if not key_text:
            continue
        if key_text in seen_keys:
            errors.append(
                ExportError(
                    workbook_path,
                    sheet.title,
                    row,
                    key_column.index,
                    f"Duplicate key: {key_text}",
                )
            )
            continue
        seen_keys.add(key_text)

        record: dict[str, Any] = {}
        for column in exported_columns:
            raw_value = sheet.cell(row, column.index).value
            try:
                record[column.name] = convert_value(raw_value, column)
            except ValueError as exc:
                errors.append(
                    ExportError(
                        workbook_path,
                        sheet.title,
                        row,
                        column.index,
                        str(exc),
                    )
                )
        rows.append(record)

    return rows


def convert_value(raw_value: Any, column: ColumnSchema) -> Any:
    if column.is_array:
        if raw_value is None or normalize_text(raw_value) == "":
            return []
        return [
            convert_scalar(part, column.data_type, column.nullable, column.name)
            for part in split_csv_like(normalize_text(raw_value))
        ]

    return convert_scalar(raw_value, column.data_type, column.nullable, column.name)


def convert_scalar(raw_value: Any, data_type: str, nullable: bool, field_name: str) -> Any:
    if raw_value is None or normalize_text(raw_value).lower() == "null":
        if nullable:
            return None
        raise ValueError(f"{field_name} is empty but is not nullable.")

    if data_type == "string":
        return normalize_text(raw_value)
    if data_type == "int":
        return parse_int(raw_value, field_name)
    if data_type == "float":
        return parse_float(raw_value, field_name)
    if data_type == "bool":
        return parse_bool(raw_value, field_name)
    if data_type == "json":
        return parse_json_cell(raw_value, field_name)

    raise ValueError(f"{field_name} uses unsupported type: {data_type}")


def parse_int(value: Any, field_name: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{field_name} must be int, got bool.")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = normalize_text(value)
    try:
        parsed = int(text)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be int: {text}") from exc
    return parsed


def parse_float(value: Any, field_name: str) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{field_name} must be float, got bool.")
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value)
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be float: {text}") from exc


def parse_bool(value: Any, field_name: str) -> bool:
    if isinstance(value, bool):
        return value
    text = normalize_text(value).lower()
    if text in {"true", "1", "yes", "y"}:
        return True
    if text in {"false", "0", "no", "n"}:
        return False
    raise ValueError(f"{field_name} must be bool: {normalize_text(value)}")


def parse_json_cell(value: Any, field_name: str) -> Any:
    if isinstance(value, (dict, list)):
        return value
    text = normalize_text(value)
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{field_name} must be valid JSON: {exc.msg}") from exc


def split_options(value: Any) -> list[str]:
    return [part.strip() for part in normalize_text(value).split(",") if part.strip()]


def split_csv_like(value: str) -> list[str]:
    return [part.strip() for part in value.split(",") if part.strip()]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def schema_error(workbook_path: Path, sheet: Any, column: int, message: str) -> ExportError:
    return ExportError(workbook_path, sheet.title, FIELD_TYPE_ROW, column, message)


if __name__ == "__main__":
    raise SystemExit(main())
