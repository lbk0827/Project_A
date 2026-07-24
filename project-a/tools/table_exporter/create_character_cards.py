#!/usr/bin/env python3
"""Create CharacterCards.xlsx with normalized card and effect authoring tables."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import openpyxl


CARD_FIELDS = [
    "CardKey",
    "Character",
    "Id",
    "DisplayName",
    "Cost",
    "CardType",
    "MotionAnimation",
    "TextTemplate",
    "Copies",
    "Keywords[]",
]
CARD_TYPES = [
    None,
    "key,string",
    "string",
    "string",
    "string",
    "int",
    "string",
    "string",
    "string",
    "int",
    "string",
]
CARD_WIDTHS = {
    "B": 24,
    "C": 12,
    "D": 22,
    "E": 18,
    "F": 8,
    "G": 12,
    "H": 18,
    "I": 72,
    "J": 8,
    "K": 18,
}

EFFECT_FIELDS = [
    "Id",
    "CardKey",
    "Character",
    "CardId",
    "Trigger",
    "Stance",
    "EffectType",
    "Target",
    "CardType",
    "Buff",
    "Percent",
    "Amount",
    "DurationTurns",
    "Scope",
    "Condition",
    "TextArgIndex",
]
EFFECT_TYPES = [
    None,
    "key,string",
    "string",
    "string",
    "string",
    "string",
    "string,null",
    "string",
    "string,null",
    "string,null",
    "string,null",
    "int,null",
    "int,null",
    "int,null",
    "string,null",
    "string,null",
    "int,null",
]
EFFECT_WIDTHS = {
    "B": 38,
    "C": 24,
    "D": 12,
    "E": 22,
    "F": 18,
    "G": 12,
    "H": 22,
    "I": 16,
    "J": 18,
    "K": 20,
    "L": 12,
    "M": 12,
    "N": 18,
    "O": 18,
    "P": 26,
    "Q": 16,
}

GLOSSARY_SECTIONS = {
    "카드 사용 관련": "카드사용",
    "이로운 효과 관련": "이로운",
    "해로운 효과 관련": "해로운",
    # Temporary compatibility with currently mojibaked glossary headers.
    "移대뱶 ?ъ슜 愿??": "카드사용",
    "?대줈???④낵 愿??": "이로운",
    "紐ъ뒪???④낵 愿??": "해로운",
}


def card_key(character: str, card_id: str) -> str:
    return f"{character}/{card_id}"


def load_generated(project_root: Path, file_name: str) -> list[dict[str, Any]]:
    path = project_root / "data" / "generated" / file_name
    return json.loads(path.read_text(encoding="utf-8"))


def build_character_cards(sheet, cards: list[dict[str, Any]]) -> None:
    sheet.append([None])
    sheet.append(["#data", *CARD_FIELDS])
    sheet.append(CARD_TYPES)
    sheet.append([None, *["data"] * len(CARD_FIELDS)])
    for card in cards:
        character = str(card.get("Character", ""))
        card_id = str(card.get("Id", ""))
        sheet.append([
            None,
            card.get("CardKey", card_key(character, card_id)),
            character,
            card_id,
            card.get("DisplayName", ""),
            card.get("Cost", 0),
            card.get("CardType", "Skill"),
            card.get("MotionAnimation", "Idle"),
            card.get("TextTemplate", card.get("Text", "")),
            card.get("Copies", 1),
            ",".join(card.get("Keywords", [])),
        ])
    for column, width in CARD_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"


def build_card_effect_rows(sheet, effect_rows: list[dict[str, Any]]) -> None:
    sheet.append([None])
    sheet.append(["#data", *EFFECT_FIELDS])
    sheet.append(EFFECT_TYPES)
    sheet.append([None, *["data"] * len(EFFECT_FIELDS)])
    for row in effect_rows:
        sheet.append([None, *[row.get(field) for field in EFFECT_FIELDS]])
    for column, width in EFFECT_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"


def parse_glossary(glossary_path: Path) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    category = ""
    for line in glossary_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        header = re.match(r"^##\s+(.*)$", stripped)
        if header:
            category = GLOSSARY_SECTIONS.get(header.group(1).strip(), "")
            continue
        if not category or not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if len(cells) < 2:
            continue
        keyword, description = cells[0], cells[1]
        if keyword in ("키워드", "?ㅼ썙??", "") or set(keyword) <= {"-", " "}:
            continue
        rows.append((category, keyword, description))
    return rows


def build_card_effects(sheet, glossary_path: Path) -> None:
    sheet.append([None])
    sheet.append(["#data", "Keyword", "Category", "Description"])
    sheet.append([None, "key,string", "string", "string"])
    sheet.append([None, "data", "data", "data"])
    for category, keyword, description in parse_glossary(glossary_path):
        sheet.append([None, keyword, category, description])
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 90
    sheet.freeze_panes = "A5"


def main() -> None:
    project_root = Path(__file__).resolve().parents[2]
    repo_root = project_root.parents[0]
    glossary_path = repo_root / "Docs" / "Card_Effects_Glossary.md"
    output_dir = project_root / "tables" / "excel"
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    build_character_cards(workbook.active, load_generated(project_root, "CharacterCards.json"))
    workbook.active.title = "CharacterCards"
    build_card_effect_rows(
        workbook.create_sheet("CardEffectRows"),
        load_generated(project_root, "CardEffectRows.json"),
    )
    build_card_effects(workbook.create_sheet("CardEffects"), glossary_path)

    workbook.save(output_dir / "CharacterCards.xlsx")


if __name__ == "__main__":
    main()
